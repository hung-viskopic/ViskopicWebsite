import hashlib
import io
import unittest
from docx import Document
from pypdf import PdfWriter
from extraction import extract

def text_pdf(value: str) -> bytes:
    stream = f'BT /F1 12 Tf 72 720 Td ({value}) Tj ET'.encode('ascii')
    objects = [
        b'<< /Type /Catalog /Pages 2 0 R >>',
        b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
        b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>',
        b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>',
        b'<< /Length ' + str(len(stream)).encode() + b' >>\nstream\n' + stream + b'\nendstream',
    ]
    document = bytearray(b'%PDF-1.4\n')
    offsets = [0]
    for number, obj in enumerate(objects, 1):
        offsets.append(len(document))
        document.extend(f'{number} 0 obj\n'.encode() + obj + b'\nendobj\n')
    xref = len(document)
    document.extend(f'xref\n0 {len(objects) + 1}\n'.encode())
    document.extend(b'0000000000 65535 f \n')
    for offset in offsets[1:]:
        document.extend(f'{offset:010d} 00000 n \n'.encode())
    document.extend(f'trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n'.encode())
    return bytes(document)

class ExtractionTests(unittest.TestCase):
    def test_preserves_unicode_and_normalises_line_endings(self):
        raw = '\ufeffFirst paragraph.\r\nCaf\u00e9 discussion.'.encode()
        self.assertEqual(extract(raw, hashlib.sha256(raw).hexdigest(), 'essay.txt'), ('First paragraph.\nCaf\u00e9 discussion.', 4))

    def test_extracts_docx_paragraphs_and_tables(self):
        document = Document()
        document.add_paragraph('First paragraph.')
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = 'Evidence'
        table.cell(0, 1).text = 'Source'
        stream = io.BytesIO()
        document.save(stream)
        raw = stream.getvalue()
        text, count = extract(raw, hashlib.sha256(raw).hexdigest(), 'essay.docx')
        self.assertEqual(text, 'First paragraph.\n\nEvidence\tSource')
        self.assertEqual(count, 4)

    def test_rejects_pdf_without_readable_text(self):
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        stream = io.BytesIO()
        writer.write(stream)
        raw = stream.getvalue()
        with self.assertRaisesRegex(ValueError, 'No readable text'):
            extract(raw, hashlib.sha256(raw).hexdigest(), 'scan.pdf')

    def test_extracts_text_pdf(self):
        raw = text_pdf('Readable PDF evidence')
        self.assertEqual(extract(raw, hashlib.sha256(raw).hexdigest(), 'essay.pdf'), ('Readable PDF evidence', 3))

    def test_rejects_corrupted_object(self):
        with self.assertRaises(ValueError):
            extract(b'altered', hashlib.sha256(b'original').hexdigest(), 'essay.txt')

if __name__ == '__main__':
    unittest.main()
