import hashlib
import io
import zipfile
from pathlib import Path

from docx import Document
from pypdf import PdfReader

PIPELINE_VERSION = 'document-v2'
MAX_PDF_PAGES = 500
MAX_DOCX_ENTRIES = 500
MAX_DOCX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024

def _extract_txt(content: bytes) -> str:
    return content.decode('utf-8-sig').replace('\r\n', '\n').replace('\r', '\n')

def _extract_pdf(content: bytes) -> str:
    reader = PdfReader(io.BytesIO(content))
    if reader.is_encrypted:
        raise ValueError('Encrypted PDF files are not supported')
    if len(reader.pages) > MAX_PDF_PAGES:
        raise ValueError(f'PDF exceeds the {MAX_PDF_PAGES}-page limit')
    pages = [(page.extract_text() or '').strip() for page in reader.pages]
    return '\n\n'.join(page for page in pages if page)

def _extract_docx(content: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        entries = archive.infolist()
        if len(entries) > MAX_DOCX_ENTRIES or sum(item.file_size for item in entries) > MAX_DOCX_UNCOMPRESSED_BYTES:
            raise ValueError('Word document expands beyond the processing limit')
        if 'word/document.xml' not in {item.filename for item in entries}:
            raise ValueError('Word document has no main document content')
    document = Document(io.BytesIO(content))
    blocks = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            value = '\t'.join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if value:
                blocks.append(value)
    return '\n\n'.join(blocks)

def extract(content: bytes, expected_hash: str, filename: str = 'document.txt') -> tuple[str, int]:
    if hashlib.sha256(content).hexdigest() != expected_hash:
        raise ValueError('Document integrity check failed')
    extension = Path(filename).suffix.lower()
    if extension == '.txt':
        text = _extract_txt(content)
    elif extension == '.pdf':
        text = _extract_pdf(content)
    elif extension == '.docx':
        text = _extract_docx(content)
    else:
        raise ValueError('Unsupported document format')
    if not text.strip():
        raise ValueError('No readable text was found in the document')
    return text, len(text.split())
