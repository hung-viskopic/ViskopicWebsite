import io
import json
import os
import time
import urllib.request
import uuid

from docx import Document

API_URL = os.getenv('API_URL', 'http://api:4300/api')
ACCESS_KEY = os.getenv('DEV_ACCESS_KEY', 'local-viskopic-development-key')

document = Document()
document.add_heading('Synthetic Word submission', level=1)
document.add_paragraph('This document contains no student data.')
table = document.add_table(rows=1, cols=2)
table.cell(0, 0).text = 'Evidence'
table.cell(0, 1).text = 'Verified'
stream = io.BytesIO()
document.save(stream)

boundary = f'----viskopic-{uuid.uuid4().hex}'
body = (
    f'--{boundary}\r\n'
    'Content-Disposition: form-data; name="file"; filename="synthetic-word-test.docx"\r\n'
    'Content-Type: application/vnd.openxmlformats-officedocument.wordprocessingml.document\r\n\r\n'
).encode() + stream.getvalue() + f'\r\n--{boundary}--\r\n'.encode()
request = urllib.request.Request(
    f'{API_URL}/submissions',
    data=body,
    method='POST',
    headers={'Content-Type': f'multipart/form-data; boundary={boundary}', 'x-access-key': ACCESS_KEY},
)
with urllib.request.urlopen(request, timeout=10) as response:
    submission_id = json.load(response)['id']

for _ in range(30):
    detail_request = urllib.request.Request(
        f'{API_URL}/submissions/{submission_id}',
        headers={'x-access-key': ACCESS_KEY},
    )
    with urllib.request.urlopen(detail_request, timeout=10) as response:
        record = json.load(response)
    if record['status'] == 'failed':
        raise RuntimeError(f'DOCX extraction failed: {record["error"]}')
    if record['status'] == 'completed':
        assert record['pipeline_version'] == 'document-v2'
        assert 'Synthetic Word submission' in record['extracted_text']
        assert 'Evidence\tVerified' in record['extracted_text']
        print(f'DOCX upload and extraction passed: {submission_id}')
        break
    time.sleep(1)
else:
    raise TimeoutError('DOCX extraction did not finish within 30 seconds')
